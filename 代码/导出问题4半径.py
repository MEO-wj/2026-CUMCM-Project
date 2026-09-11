#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从第四问未舍入结果导出与result4.xlsx同刻的实际半径，不重复求解。"""
import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--result-dir', type=Path,
                        default=Path(__file__).resolve().parent.parent / '题目要求结果数据')
    args = parser.parse_args()
    with np.load(args.result_dir / '复算依据' / 'result4_未舍入.npz', allow_pickle=False) as raw:
        times = raw['time_s']
        radii_cm = raw['R_m'] * 100
    if (times.ndim != 1 or radii_cm.shape != times.shape or len(times) == 0
            or not np.isfinite(times).all() or not np.isfinite(radii_cm).all()
            or times[0] != 0 or np.any(np.diff(times) <= 0) or np.any(radii_cm <= 0)):
        raise ValueError('第四问的时间或半径数据不完整。')

    rows = [[float(format(value, '.4f')) for value in (t, t / 3600, radius)]
            for t, radius in zip(times, radii_cm)]
    original = load_workbook(args.result_dir / 'result4.xlsx', read_only=True, data_only=True)
    try:
        source_times = [row[0] for row in original['Sheet1'].iter_rows(min_row=2, max_col=1, values_only=True)]
    finally:
        original.close()
    if source_times != [row[0] for row in rows]:
        raise ValueError('未舍入结果的时间与result4.xlsx不一致，停止导出。')

    book = Workbook()
    sheet = book.active
    sheet.title = '实际半径'
    sheet.append(['时间/s', '时间/h', '实际半径/cm'])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(name='Songti SC', size=11, bold=True, color='000000')
        cell.fill = PatternFill('solid', fgColor='F2F2F2')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(bottom=Side(style='thin', color='808080'))
    number_font = Font(name='Times New Roman', size=11, color='000000')
    number_alignment = Alignment(horizontal='right', vertical='center')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = '0.0000'
            cell.font = number_font
            cell.alignment = number_alignment
    sheet.row_dimensions[1].height = 28
    sheet.sheet_format.defaultRowHeight = 20
    for column in ('A', 'B', 'C'):
        sheet.column_dimensions[column].width = 22
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:C{sheet.max_row}'
    sheet.sheet_view.showGridLines = False
    output = args.result_dir / '问题4半径' / '问题4半径.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    print(json.dumps({'output': str(output), 'data_rows': len(rows),
                      'end_time_s': rows[-1][0], 'end_time_h': rows[-1][1],
                      'end_radius_cm': rows[-1][2]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
