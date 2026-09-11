#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按原题格式，从四份完整结果中摘取论文表1～表6；不重复求解或绘图。"""
import argparse
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def extract_table(path, sheet_name, requested_times=None, moving_surface=False):
    """只读取规定时刻；后两表每6小时取一行，并保留源文件实际末行。"""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        source = book[sheet_name]
        iterator = source.iter_rows(values_only=True)
        headers = next(iterator)
        selected = []
        last = None
        for row in iterator:
            if row[0] is None:
                continue
            time_s = row[0]
            if requested_times is not None and time_s > max(requested_times):
                break
            last = row
            wanted = (time_s in requested_times if requested_times is not None
                      else time_s > 0 and abs(time_s % 21600) < 1e-7)
            if wanted:
                selected.append(row)
        if requested_times is not None:
            if [row[0] for row in selected] != list(requested_times):
                raise ValueError(f'{path.name}/{sheet_name} 缺少题目规定时刻。')
        elif last is None:
            raise ValueError(f'{path.name}/{sheet_name} 没有结果数据。')
        elif not selected or selected[-1][0] != last[0]:
            selected.append(last)

        columns = [j for j, radius in enumerate(headers) if j > 0
                   and isinstance(radius, (int, float))
                   and abs(radius * 2 - round(radius * 2)) < 1e-9]
        if moving_surface:
            # 不列出在所有规定时刻都已位于药材外部的固定测点。
            columns = [j for j in columns if any(row[j] is not None for row in selected)]
            columns.append(headers.index('药材表面'))
        elif [headers[j] for j in columns] != [0, .5, 1, 1.5, 2]:
            raise ValueError(f'{path.name}/{sheet_name} 的径向表头与原题不符。')

        rows = [[row[0], *[row[j] for j in columns]] for row in selected]
        if any(value is not None and not math.isfinite(value) for row in rows for value in row):
            raise ValueError(f'{path.name}/{sheet_name} 包含非有限数值。')
        if not moving_surface and any(value is None for row in rows for value in row):
            raise ValueError(f'{path.name}/{sheet_name} 包含缺失测点。')
        return [headers[j] for j in columns], rows
    finally:
        book.close()


def add_table(book, number, title, radii, source_rows, time_unit):
    sheet = book.create_sheet(f'表{number}')
    last_column = get_column_letter(len(radii) + 1)
    last_row = len(source_rows) + 4
    text_font = Font(name='Songti SC', size=11, color='000000')
    number_font = Font(name='Times New Roman', size=11, color='000000')
    thin = Side(style='thin', color='000000')

    sheet['A1'] = f'表{number}  {title}'
    sheet.merge_cells(f'A1:{last_column}1')
    sheet['A1'].font = Font(name='Songti SC', size=12, bold=True, color='000000')
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 7
    sheet['A3'] = f'时间/{time_unit}'
    sheet['B3'] = '到药材中心的距离/cm'
    for column, radius in enumerate(radii, 2):
        sheet.cell(4, column, radius)
    for row_index, values in enumerate(source_rows, 5):
        divisor = 1 if time_unit == 's' else 3600
        sheet.cell(row_index, 1, float(format(values[0] / divisor, '.4f')))
        for column, value in enumerate(values[1:], 2):
            sheet.cell(row_index, column, value)

    for row in sheet.iter_rows(min_row=3, max_row=last_row, max_col=len(radii) + 1):
        for cell in row:
            cell.font = number_font if isinstance(cell.value, (int, float)) else text_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.0000'
    sheet.merge_cells('A3:A4')
    sheet.merge_cells(f'B3:{last_column}3')
    for row_index in range(3, last_row + 1):
        sheet.row_dimensions[row_index].height = 23
    if number >= 5:
        # 仍以数值存储小时数；显示时同时保留原题的末行标识。
        sheet.cell(last_row, 1).number_format = '"烘干结束时间\n"0.0000'
        sheet.cell(last_row, 1).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.cell(last_row, 1).font = text_font
        sheet.row_dimensions[last_row].height = 39
    sheet.column_dimensions['A'].width = 22
    for column in range(2, len(radii) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    sheet.print_area = f'A1:{last_column}{last_row}'
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=.5, right=.5, top=.5, bottom=.5, header=0, footer=0)
    sheet.print_options.horizontalCentered = True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--result-dir', type=Path,
                        default=Path(__file__).resolve().parent.parent / '题目要求结果数据')
    args = parser.parse_args()
    specs = [
        (1, '30分钟内药材的温度', 1, '温度', [100, 300, 600, 900, 1200, 1500, 1800], 's'),
        (2, '30分钟内药材的水分浓度', 1, '水分浓度', [100, 300, 600, 900, 1200, 1500, 1800], 's'),
        (3, '3小时内药材的温度', 2, '温度', [1800, 3600, 5400, 7200, 9000, 10800], 'h'),
        (4, '3小时内药材的水分浓度', 2, '水分浓度', [1800, 3600, 5400, 7200, 9000, 10800], 'h'),
        (5, '药材烘干过程的水分浓度', 3, 'Sheet1', None, 'h'),
        (6, '药材烘干过程的水分浓度', 4, 'Sheet1', None, 'h'),
    ]
    book = Workbook()
    book.remove(book.active)
    summary = []
    for number, title, source_number, sheet_name, times, unit in specs:
        radii, rows = extract_table(args.result_dir / f'result{source_number}.xlsx',
                                   sheet_name, times, moving_surface=number == 6)
        add_table(book, number, title, radii, rows, unit)
        summary.append({'table': number, 'rows': len(rows), 'radii_cm': radii,
                        'result_values': sum(value is not None for row in rows for value in row[1:])})
    output = args.result_dir / '表1到表6' / '表1到表6.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    print(json.dumps({'output': str(output), 'tables': summary}, ensure_ascii=False))


if __name__ == '__main__':
    main()
