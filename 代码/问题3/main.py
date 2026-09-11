#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""问题3：药材烘干模型的独立求解与完整结果导出。

输入：代码/附件/附件1.xlsx、附件2.xlsx（原题数据）。
输出：相邻“题目要求结果数据”中的result3.xlsx及未舍入复算依据。
默认采用2560/5120两网格整场Richardson处理，动态求临界时间；不画图。
依赖：NumPy、SciPy、openpyxl。全部模型参数和结果生成方法均在本文件。
"""
import os
os.environ['OPENBLAS_NUM_THREADS'] = '1'
os.environ['OMP_NUM_THREADS'] = '1'

QUESTION = 3

import sys
import time
import numpy as np
import scipy
from scipy.integrate import BDF
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.sparse import diags

SOLVER_VERSION = 'radial-bdf-pchip-whole-field-richardson-20260911'


def output_geometry(question, end_s, radius_data):
    spacing = 60. if question == 4 else 1.
    times = np.arange(0., np.floor(end_s / spacing) * spacing + .01, spacing)
    if times[-1] < end_s:
        times = np.r_[times, end_s]
    if question == 4:
        radius = PchipInterpolator(radius_data[:, 0], radius_data[:, 1] * .01)
        R = radius(times)
        physical = np.r_[np.arange(20) / 10., np.nan]
        domain = np.c_[physical[None, :-1] * .01 <= R[:, None] + 1e-14,
                       np.ones(len(times), dtype=bool)]
        headers = np.array([f'{v:.1f} cm' for v in physical[:-1]] + ['actual_surface'])
    else:
        R = np.full(len(times), .02)
        physical = np.arange(21) / 10.
        domain = np.ones((len(times), 21), dtype=bool)
        headers = np.array([f'{v:.1f} cm' for v in physical])
    return times, physical, headers, R, domain


def integrate_direct(question, cells, environment, radius_data, progress=None):
    """Stream one direct grid, finishing the block that brackets its event."""
    end_limit = 1800. if question == 1 else (float(radius_data[-1, 0]) if question == 4 else 96. * 3600.)
    times, radii, headers, R_output, domain = output_geometry(question, end_limit, radius_data)
    C_out = np.full((len(times), 21), np.nan)
    T_out = np.full((len(times), 21), np.nan)
    C_out[0], T_out[0] = 2.55, 28.
    xi = 1 - (1 - np.linspace(0, 1, cells + 1)) ** 1.5
    n = cells + 1
    state = np.tile([2.55, 28.], n)
    faces = np.r_[0., (xi[1:] + xi[:-1]) / 2, 1.]
    volumes = np.diff(faces ** 2) / 2
    dx = np.diff(xi)
    gauss, weights = np.polynomial.legendre.leggauss(8)
    gauss = (gauss[:, None] + 1) / 2
    weights = weights / 2
    raw = np.asarray(environment, dtype=float)
    plateau_rows = raw[(raw[:, 0] >= 10800.) & (raw[:, 0] <= 14400.), 1:]
    if plateau_rows.shape != (61, 2):
        raise ValueError('Attachment 1 must contain all 61 records from 10800 to 14400 s.')
    plateau = plateau_rows.mean(axis=0)
    radius = PchipInterpolator(radius_data[:, 0], radius_data[:, 1] * .01)
    sparse = diags([np.ones(2 * n)] * 7, range(-3, 4), shape=(2 * n, 2 * n), format='csc')
    knots = np.r_[0., raw[(raw[:, 0] > 0) & (raw[:, 0] < end_limit), 0], end_limit]
    if question == 4:
        knots = np.r_[knots, radius_data[(radius_data[:, 0] > 0) & (radius_data[:, 0] < end_limit), 0]]
    elif question != 1:
        knots = np.r_[knots, np.arange(16200., end_limit, 1800.)]
    knots = np.unique(knots)
    A, a, E = (7e-9, .89, 0.) if question == 1 else ((4.2e-4, .30, 3850.) if question == 4 else (2.4e-3, .45, 3850.))
    rtol = 2e-12
    event_time = None
    index = 1
    nfev = njev = nlu = 0
    previous_block = []
    current_block = []
    tic = last_progress = time.perf_counter()
    for begin, finish in zip(knots[:-1], knots[1:]):
        previous_block, current_block = current_block, []
        if begin >= 14400.:
            va, slope = plateau, np.zeros(2)
        else:
            va = np.array([np.interp(begin, raw[:, 0], raw[:, j]) for j in (1, 2)])
            vb = np.array([np.interp(finish, raw[:, 0], raw[:, j]) for j in (1, 2)])
            slope = (vb - va) / (finish - begin)

        def rhs(t, y):
            ct = y.reshape(n, 2)
            c, temperature = ct[:, 0], ct[:, 1]
            R = float(radius(t)) if question == 4 else .02
            tair, cair = va + (t - begin) * slope
            dc = np.diff(c)
            cq = c[:-1][None, :] + gauss * dc[None, :]
            effective = weights @ np.exp(-a / cq)
            thermal = np.exp(-E / (temperature + 273.15))
            df = A * effective * (thermal[:-1] + thermal[1:]) / 2
            if question == 1:
                kf, capacity = .36, 820 * 2600.
            elif question == 4:
                k = .12 + .20 * c / (1 + c)
                kf = (k[:-1] + k[1:]) / 2
                capacity = (760 + 90 * c) * (1850 + 2150 * c / (1 + c))
            else:
                k = .21 + .38 * c / (1 + c)
                kf = (k[:-1] + k[1:]) / 2
                capacity = (650 + 128 * c) * (1450 + 2736 * c / (1 + c))
            qc = np.r_[0., faces[1:-1] * df * dc / dx, R * 8e-7 * (cair - c[-1])]
            qt = np.r_[0., faces[1:-1] * kf * np.diff(temperature) / dx, R * 25 * (tair - temperature[-1])]
            rate = np.empty_like(ct)
            rate[:, 0] = np.diff(qc) / (R * R * volumes)
            rate[:, 1] = np.diff(qt) / (R * R * volumes * capacity)
            return rate.ravel()

        solver = BDF(rhs, begin, state, finish, rtol=rtol,
                     atol=np.tile([1e-14, 1e-13], n), max_step=60., jac_sparsity=sparse)
        while solver.status == 'running':
            t_before = solver.t
            maximum_before = float(solver.y[0::2].max())
            message = solver.step()
            if solver.status == 'failed':
                raise RuntimeError(message)
            dense = solver.dense_output()
            current_block.append(dense)
            stop_index = int(np.searchsorted(times, solver.t, side='right'))
            if stop_index > index:
                chosen = times[index:stop_index]
                native = dense(chosen).reshape(n, 2, -1)
                if question == 4:
                    for j in range(len(chosen)):
                        targets = np.r_[np.arange(20) / 1000. / R_output[index + j], 1.]
                        values = PchipInterpolator(xi, native[:, :, j], axis=0)(np.minimum(targets, 1.))
                        values[~domain[index + j]] = np.nan
                        C_out[index + j], T_out[index + j] = values.T
                else:
                    values = PchipInterpolator(xi, native, axis=0)(radii / 2.).transpose(2, 0, 1)
                    C_out[index:stop_index], T_out[index:stop_index] = values[:, :, 0], values[:, :, 1]
                index = stop_index
            if event_time is None and maximum_before > .15 >= float(solver.y[0::2].max()):
                event_time = brentq(lambda t: float(dense(t)[0::2].max()) - .15,
                                   t_before, solver.t, xtol=1e-8)
            now = time.perf_counter()
            if progress is not None and now - last_progress >= 20.:
                progress({'question': question, 'grid': cells, 'time_s': solver.t, 'elapsed_s': now - tic})
                last_progress = now
        # Match solve_ivp(t_eval=[..., finish]): propagate the same dense endpoint
        # that supplies the output at each original input breakpoint.
        state = current_block[-1](finish)
        nfev += solver.nfev
        njev += solver.njev
        nlu += solver.nlu
        if event_time is not None and finish > event_time + 1.:
            break
    if question != 1 and event_time is None:
        raise RuntimeError('No drying threshold crossing within the declared search horizon.')
    blocks = previous_block + current_block
    block_end = np.array([d.t for d in blocks])
    def evaluate(t):
        j = min(int(np.searchsorted(block_end, t, side='left')), len(blocks) - 1)
        if not blocks[j].t_old - 1e-8 <= t <= blocks[j].t + 1e-8:
            raise ValueError('Requested time is outside the retained final integration blocks.')
        return blocks[j](t)
    return {'time_s': times[:index], 'C': C_out[:index].copy(), 'T': T_out[:index].copy(),
            'xi': xi, 'state': state.reshape(n, 2), 'evaluate': evaluate,
            'retained_start_s': float(blocks[0].t_old), 'end_s': float(finish),
            'metadata': {'cells': cells, 'grading': 1.5, 'rtol': rtol, 'atol_C': 1e-14,
                         'atol_T': 1e-13, 'max_step_s': 60., 'nfev': nfev, 'njev': njev,
                         'nlu': nlu, 'event_time_s': event_time,
                         'elapsed_s': time.perf_counter() - tic,
                         'integration_boundaries_s': knots[knots <= finish].tolist()}}


def maximum_reconstruction(xfine, fine, xcoarse, coarse):
    """Find the maximum of all pieces of (4 PCHIP_fine - PCHIP_coarse)/3."""
    p, q = PchipInterpolator(xfine, fine), PchipInterpolator(xcoarse, coarse)
    left = xfine[:-1]
    j = np.minimum(np.searchsorted(xcoarse, left, side='right') - 1, len(xcoarse) - 2)
    delta = left - xcoarse[j]
    a, b, c, d = q.c[:, j]
    shifted = np.array([a, 3 * a * delta + b, 3 * a * delta ** 2 + 2 * b * delta + c,
                        a * delta ** 3 + b * delta ** 2 + c * delta + d])
    coeff = (4 * p.c - shifted) / 3.
    width = np.diff(xfine)
    maximum = float(np.max((4 * fine - q(xfine)) / 3.))
    A, B, C = 3 * coeff[0], 2 * coeff[1], coeff[2]
    discriminant = B * B - 4 * A * C
    valid = (discriminant >= 0) & (np.abs(A) > 1e-30)
    roots = []
    for sign in (-1., 1.):
        r = np.full_like(left, np.nan)
        r[valid] = (-B[valid] + sign * np.sqrt(discriminant[valid])) / (2 * A[valid])
        roots.append(r)
    linear = (np.abs(A) <= 1e-30) & (np.abs(B) > 1e-30)
    r = np.full_like(left, np.nan)
    r[linear] = -C[linear] / B[linear]
    roots.append(r)
    for r in roots:
        inside = (r > 0) & (r < width)
        if np.any(inside):
            x = r[inside]
            a, b, c, d = coeff[:, inside]
            maximum = max(maximum, float(np.max(((a * x + b) * x + c) * x + d)))
    return maximum


def solve_problem(question, environment, radius_data, coarse_cells=2560, progress=None):
    """Return unrounded whole-field results; Q2 and Q3 share one trajectory rule."""
    if question not in (1, 2, 3, 4):
        raise ValueError('question must be 1, 2, 3, or 4')
    model = 2 if question == 3 else question
    coarse = integrate_direct(model, coarse_cells, environment, radius_data, progress)
    fine = integrate_direct(model, 2 * coarse_cells, environment, radius_data, progress)
    critical = None
    if model == 1:
        end = 1800.
    else:
        def residual(t):
            return maximum_reconstruction(fine['xi'], fine['evaluate'](t)[0::2],
                                          coarse['xi'], coarse['evaluate'](t)[0::2]) - .15
        a = max(coarse['retained_start_s'], fine['retained_start_s'])
        b = min(coarse['end_s'], fine['end_s'])
        if not residual(a) > 0 > residual(b):
            raise RuntimeError('The common retained blocks do not bracket the extrapolated threshold.')
        critical = brentq(residual, a, b, xtol=1e-8)
        end = (round(critical / 3600., 4) + .0001) * 3600.
        if not critical < end <= b or residual(end) >= 0:
            raise RuntimeError('The dynamically selected terminal state does not strictly satisfy C < 0.15.')
    times, radii, headers, R, domain = output_geometry(model, end, radius_data)
    direct_outputs = []
    endpoint_states = []
    for grid in (coarse, fine):
        rows = np.searchsorted(grid['time_s'], times)
        fields = {'C': grid['C'][rows].copy(), 'T': grid['T'][rows].copy()}
        final_state = grid['state'] if model == 1 else grid['evaluate'](end).reshape(-1, 2)
        endpoint_states.append(final_state)
        if model != 1:
            targets = np.r_[radii[:-1] * .01 / R[-1], 1.] if model == 4 else radii / 2.
            values = PchipInterpolator(grid['xi'], final_state, axis=0)(np.minimum(targets, 1.))
            values[~domain[-1]] = np.nan
            fields['C'][-1], fields['T'][-1] = values.T
        direct_outputs.append(fields)
    result = {'time_s': times, 'radius_cm': radii, 'radius_headers': headers, 'R_m': R, 'domain_mask': domain}
    for field in ('C', 'T'):
        result[field] = (4 * direct_outputs[1][field] - direct_outputs[0][field]) / 3.
        if not np.isfinite(result[field][domain]).all() or not np.isnan(result[field][~domain]).all():
            raise RuntimeError('Nonfinite values or an incorrect outside-domain mask were produced.')
    native = (4 * endpoint_states[1] - PchipInterpolator(coarse['xi'], endpoint_states[0], axis=0)(fine['xi'])) / 3.
    result.update(final_xi=fine['xi'], final_C=native[:, 0], final_T=native[:, 1])
    if result['C'][domain].min() <= 0 or result['C'][domain].max() > 2.55 + 1e-8:
        raise RuntimeError('The complete extrapolated moisture field violates its physical range.')
    if result['T'][domain].min() < 28. - 1e-8 or result['T'][domain].max() > np.max(environment[:, 1]) + 1e-8:
        raise RuntimeError('The complete extrapolated temperature field violates its physical range.')
    if question == 3:
        rows = np.r_[np.arange(0, len(times) - 1, 60), len(times) - 1]
        for key in ('time_s', 'R_m', 'domain_mask', 'C', 'T'):
            result[key] = result[key][rows]
        for fields in direct_outputs:
            for key in ('C', 'T'):
                fields[key] = fields[key][rows]
    result['metadata'] = {
        'solver_version': SOLVER_VERSION, 'question': question, 'model_question': model,
        'python': sys.version, 'numpy': np.__version__, 'scipy': scipy.__version__,
        'method': 'BDF with conservative graded radial nodes and spatial PCHIP',
        'extrapolation': '(4*fine-coarse)/3 applied uniformly to every C/T output value',
        'grids': [coarse['metadata'], fine['metadata']],
        'critical_time_s': critical, 'critical_time_h': None if critical is None else critical / 3600.,
        'stop_time_s': end, 'stop_time_h': end / 3600.,
        'stop_rule': 'Q1: 1800 s; Q2/Q3/Q4: (round(critical_time_h,4)+0.0001)*3600',
        'terminal_max_C_full_reconstruction': float(native[:, 0].max()) if model == 1 else residual(end) + .15,
        'plateau': environment[(environment[:, 0] >= 10800.) & (environment[:, 0] <= 14400.), 1:].mean(axis=0).tolist(),
        'physical_assumptions': '1D radial; no latent heat; h=25, hm=8e-7; Q4 affine material motion without added drift',
        'precision_scope': 'Numerical convergence evidence does not certify each fourth decimal mathematically.'}
    result['direct_coarse'] = direct_outputs[0]
    result['direct_fine'] = direct_outputs[1]
    return result


from pathlib import Path
import argparse
import json
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, NamedStyle
from openpyxl.utils import get_column_letter


def read_attachment(path, expected_columns):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.worksheets[0].values)
    finally:
        workbook.close()
    data = np.asarray([row[:expected_columns] for row in rows[1:]
                       if any(value is not None for value in row[:expected_columns])], dtype=float)
    if data.ndim != 2 or data.shape[1] != expected_columns or not np.isfinite(data).all():
        raise ValueError(f'附件数据格式或数值不完整：{path.name}')
    if data[0, 0] != 0 or not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError(f'时间必须从0开始并严格递增：{path.name}')
    return data


def progress(record):
    print(f"问题{QUESTION} 网格{record['grid']}："
          f"已计算至 {record['time_s'] / 3600:.4f} h",
          file=sys.stderr, flush=True)


def numeric_cell(sheet, value, header=False):
    if not np.isfinite(value):
        return None
    cell = WriteOnlyCell(sheet, value=float(format(float(value), '.4f')))
    cell.style = 'heading_number' if header else 'four_decimals'
    return cell


def write_excel(path, result):
    workbook = Workbook(write_only=True)
    workbook.add_named_style(NamedStyle(
        name='four_decimals', number_format='0.0000',
        font=Font(name='Arial', size=10),
        alignment=Alignment(horizontal='right', vertical='center')))
    workbook.add_named_style(NamedStyle(
        name='heading_number', number_format='0.0000',
        font=Font(name='Arial', size=10, bold=True),
        alignment=Alignment(horizontal='center', vertical='center')))
    names = [('温度', 'T'), ('水分浓度', 'C')] if QUESTION in (1, 2) else [('Sheet1', 'C')]
    count = len(result['time_s'])
    for name, field in names:
        sheet = workbook.create_sheet(name)
        sheet.freeze_panes = 'B2'
        sheet.column_dimensions['A'].width = 32
        for column in range(2, 23):
            sheet.column_dimensions[get_column_letter(column)].width = 11
        sheet.row_dimensions[1].height = 30
        first = WriteOnlyCell(sheet, value=r'时间\到药材中心的距离')
        first.font = Font(name='Arial', size=10, bold=True)
        first.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        headers = [first]
        for index, radius in enumerate(result['radius_cm']):
            if QUESTION == 4 and index == 20:
                cell = WriteOnlyCell(sheet, value='药材表面')
                cell.style = 'heading_number'
                headers.append(cell)
            else:
                headers.append(numeric_cell(sheet, radius, header=True))
        sheet.append(headers)
        for index, (t, values) in enumerate(zip(result['time_s'], result[field]), start=1):
            sheet.append([numeric_cell(sheet, t)] +
                         [numeric_cell(sheet, value) for value in values])
            if index % 40000 == 0:
                print(f'{path.name} / {name}：{index}/{count} 行',
                      file=sys.stderr, flush=True)
    pending = path.with_name('.' + path.stem + '.pending.xlsx')
    workbook.save(pending)
    pending.replace(path)


def write_evidence(folder, result):
    folder.mkdir(parents=True, exist_ok=True)
    stem = f'result{QUESTION}'
    arrays = {key: value for key, value in result.items() if isinstance(value, np.ndarray)}
    for kind in ('direct_coarse', 'direct_fine'):
        for field in ('C', 'T'):
            arrays[kind + '_' + field] = result[kind][field]
    raw_path = folder / (stem + '_未舍入.npz')
    temporary = folder / ('.' + stem + '_未舍入.pending.npz')
    np.savez_compressed(temporary, **arrays)
    temporary.replace(raw_path)
    metadata = dict(result['metadata'])
    metadata.update({
        'input_files': ['附件1.xlsx', '附件2.xlsx'],
        'output_workbook': stem + '.xlsx',
        'raw_output': raw_path.name,
        'data_rows': len(result['time_s']),
        'radius_columns': 21,
        'temperature_unit': 'degC',
        'moisture_unit': 'kg_water/kg_dry_solid',
        'time_unit': 's',
        'radius_unit': 'cm',
        'excel_rounding': "float(format(value, '.4f'))",
        'excel_number_format': '0.0000',
        'excel_blank_rule': 'Only fixed Q4 radius positions outside R(t) are empty.',
        'threshold_check': 'Use the unrounded full radial reconstruction, not rounded Excel values.',
        'openpyxl': __import__('openpyxl').__version__,
    })
    record = folder / (stem + '_生成记录.json')
    temporary = folder / ('.' + stem + '_生成记录.pending.json')
    temporary.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(record)


def main():
    code_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=f'问题{QUESTION}：完整求解并导出四位小数结果')
    parser.add_argument('--input-dir', type=Path, default=code_root / '附件')
    parser.add_argument('--output-dir', type=Path, default=code_root.parent / '题目要求结果数据')
    parser.add_argument('--coarse-cells', type=int, default=2560,
                        help='较粗网格区间数；较细网格为其两倍，正式默认2560/5120')
    args = parser.parse_args()
    if args.coarse_cells < 16:
        parser.error('--coarse-cells 至少为16；改变默认值会改变数值误差')
    environment = read_attachment(args.input_dir / '附件1.xlsx', 3)
    radius_data = read_attachment(args.input_dir / '附件2.xlsx', 2)
    if np.any(radius_data[:, 1] <= 0):
        raise ValueError('附件2半径必须为正')
    result = solve_problem(QUESTION, environment, radius_data,
                           coarse_cells=args.coarse_cells, progress=progress)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_evidence(args.output_dir / '复算依据', result)
    output = args.output_dir / f'result{QUESTION}.xlsx'
    write_excel(output, result)
    metadata = result['metadata']
    summary = {
        'question': QUESTION,
        'workbook': str(output),
        'data_rows': len(result['time_s']),
        'critical_time_h': None if metadata['critical_time_h'] is None
                           else format(metadata['critical_time_h'], '.4f'),
        'stop_time_h': format(metadata['stop_time_h'], '.4f'),
        'stop_time_s': format(metadata['stop_time_s'], '.4f'),
        'all_excel_numbers': '0.0000',
    }
    print(json.dumps({'results': [summary]}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
